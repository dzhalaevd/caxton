from collections.abc import Iterator, Mapping
from io import BytesIO

import pytest
from openpyxl import load_workbook

from caxton import (  # noqa: WPS347
    CaxtonTypeError,
    CaxtonValueError,
    DataSourceConsumedError,
    DocumentTheme,
    InvalidOperationError,
    Style,
    StyleSheet,
    UnsupportedFeatureError,
    ValidationError,
    chart,
    col,
    compose,
    decimal,
    image,
    matrix,
    render,
    sheet,
    sheet_ref,
    spacer,
    spreadsheet,
    stack,
    table,
    table_ref,
    text,
    title,
    validate,
    when,
    write,
)  # noqa: WPS347
from caxton.core.models import (
    CellReference,
    Chart,
    Formula,
    FormulaBinary,
    RangeReference,
    SpreadsheetDocument,
    Stack,
    TemplateSpecification,
)
from caxton.testing import inspect_artifact


class _ObservableRows:
    def __init__(self) -> None:
        self.iterations = 0

    def __iter__(self) -> Iterator[dict[str, int]]:
        self.iterations += 1
        yield {"value": 1}


class _UnknownFormula(Formula):
    __slots__ = ()


class _RepeatedKeys(Mapping[str, SpreadsheetDocument]):
    def __init__(self, names: tuple[str, ...], document: SpreadsheetDocument) -> None:
        self._names = names
        self._document = document

    def __getitem__(self, key: str) -> SpreadsheetDocument:
        if key not in self._names:
            raise KeyError(key)
        return self._document

    def __iter__(self) -> Iterator[str]:
        return iter(self._names)

    def __len__(self) -> int:
        return len(self._names)


def test_compose_names_worksheets_in_section_order() -> None:
    params = spreadsheet(sheet("Summary"), sheet("Details"))
    vulnerabilities = spreadsheet(sheet("Summary"))

    result = compose(
        {
            "Params": params,
            "Vulnerabilities": vulnerabilities,
        },
    )

    assert [worksheet.name for worksheet in result.worksheets] == [
        "Params - Summary",
        "Params - Details",
        "Vulnerabilities - Summary",
    ]
    assert [worksheet.blocks for worksheet in result.worksheets] == [(), (), ()]
    assert params.worksheets[0].name == "Summary"
    assert vulnerabilities.worksheets[0].name == "Summary"


@pytest.mark.parametrize("backend", ["xlsxwriter", "openpyxl"])
def test_xlsx_render_rejects_overlong_composed_worksheet_name_before_output(
    backend: str,
) -> None:
    section = "A" * 20
    local_name = "B" * 20
    worksheet_name = f"{section} - {local_name}"
    target = BytesIO()

    with pytest.raises(UnsupportedFeatureError) as captured:
        write(
            compose({section: spreadsheet(sheet(local_name))}),
            target,
            backend=backend,
        )

    assert target.getvalue() == b""
    assert captured.value.context == {
        "constraint": "maximum_length",
        "length": len(worksheet_name),
        "maximum": 31,
        "worksheet": worksheet_name,
    }


@pytest.mark.parametrize("backend", ["xlsxwriter", "openpyxl"])
@pytest.mark.parametrize(
    ("section", "local_name", "constraint", "detail"),
    [
        ("Section", "Bad/Name", "invalid_character", {"character": "/"}),
        ("'Section", "Name", "surrounding_apostrophe", {}),
        ("Section", "Name'", "surrounding_apostrophe", {}),
    ],
)
def test_xlsx_render_rejects_other_invalid_composed_worksheet_names(
    backend: str,
    section: str,
    local_name: str,
    constraint: str,
    detail: dict[str, object],
) -> None:
    worksheet_name = f"{section} - {local_name}"
    target = BytesIO()

    with pytest.raises(UnsupportedFeatureError) as captured:
        write(
            compose({section: spreadsheet(sheet(local_name))}),
            target,
            backend=backend,
        )

    assert target.getvalue() == b""
    assert captured.value.context == {
        "constraint": constraint,
        "worksheet": worksheet_name,
        **detail,
    }


def test_compose_rejects_sections_without_worksheets() -> None:
    with pytest.raises(CaxtonValueError, match="at least one worksheet") as captured:
        compose({"Empty": spreadsheet()})

    assert captured.value.context == {"reason": "empty_composition"}


def test_compose_skips_zero_worksheet_children() -> None:
    result = compose(
        {
            "Empty": spreadsheet(),
            "Data": spreadsheet(sheet("Summary")),
        },
    )

    assert tuple(worksheet.name for worksheet in result.worksheets) == (
        "Data - Summary",
    )


def test_zero_worksheet_child_contributes_no_settings() -> None:
    live_style = Style(fill="#D9EAF7")
    live_theme = DocumentTheme(default=Style(fill="#FFFFFF"))
    empty = spreadsheet(
        metadata={"owner": "empty"},
        styles={"heading": Style(fill="#FCE8E6")},
        theme=DocumentTheme(default=Style(fill="#000000")),
    )
    live = spreadsheet(
        sheet("Summary"),
        metadata={"owner": "live"},
        styles={"heading": live_style},
        theme=live_theme,
    )

    result = compose({"Empty": empty, "Live": live})

    assert result.metadata == {"owner": "live"}
    assert result.styles["heading"] is live_style
    assert result.theme is live_theme


@pytest.mark.parametrize(
    ("sections", "error", "message"),
    [
        ([], CaxtonTypeError, "Sections must be a mapping"),
        ({1: spreadsheet(sheet("Data"))}, CaxtonTypeError, "Section name"),
        ({" ": spreadsheet(sheet("Data"))}, CaxtonValueError, "cannot be empty"),
        ({"Data": object()}, CaxtonTypeError, "spreadsheet document"),
    ],
)
def test_compose_rejects_invalid_sections(
    sections: object,
    error: type[Exception],
    message: str,
) -> None:
    with pytest.raises(error, match=message):
        compose(sections)  # type: ignore[arg-type]


def test_compose_reports_invalid_child_document_section_context() -> None:
    with pytest.raises(CaxtonTypeError) as captured:
        compose({"Data": object()})  # type: ignore[dict-item]

    assert captured.value.path == "sections['Data']"
    assert captured.value.context == {"section": "Data", "type": "object"}


def test_compose_rejects_case_insensitive_section_duplicates() -> None:
    sections = {
        "Params": spreadsheet(sheet("Summary")),
        "params": spreadsheet(sheet("Details")),
    }

    with pytest.raises(InvalidOperationError) as captured:
        compose(sections)

    assert captured.value.context == {
        "reason": "duplicate_section",
        "section": "params",
        "conflicts_with": "Params",
    }


def test_compose_rejects_repeated_keys_from_a_custom_mapping() -> None:
    document = spreadsheet(sheet("Summary"))
    sections = _RepeatedKeys(("Params", "Params"), document)

    with pytest.raises(InvalidOperationError) as captured:
        compose(sections)

    assert captured.value.context == {
        "reason": "duplicate_section",
        "section": "Params",
        "conflicts_with": "Params",
    }


def test_compose_rejects_colliding_output_worksheet_names() -> None:
    sections = {
        "A": spreadsheet(sheet("B - C")),
        "A - B": spreadsheet(sheet("C")),
    }

    with pytest.raises(InvalidOperationError) as captured:
        compose(sections)

    assert captured.value.context == {
        "reason": "worksheet_name_conflict",
        "worksheet": "A - B - C",
        "section": "A - B",
        "local_worksheet": "C",
        "conflicts_with_section": "A",
        "conflicts_with_worksheet": "B - C",
    }


def test_compose_rejects_casefolded_output_worksheet_name_collision() -> None:
    sections = {
        "A": spreadsheet(sheet("B - C")),
        "a - b": spreadsheet(sheet("c")),
    }

    with pytest.raises(InvalidOperationError) as captured:
        compose(sections)

    assert captured.value.context == {
        "reason": "worksheet_name_conflict",
        "worksheet": "a - b - c",
        "section": "a - b",
        "local_worksheet": "c",
        "conflicts_with_section": "A",
        "conflicts_with_worksheet": "B - C",
    }


def test_compose_rebases_section_local_formula_references() -> None:
    rates = table(
        source=[{"rate": 2}],
        columns=(decimal(source="rate"),),
        name="rates",
    )
    amount = decimal(
        id="amount",
        formula=(sheet_ref("Rates").table("rates").column("rate").cell(0) + 1),
    )
    sales = table(
        source=[{}],
        columns=(amount,),
        name="sales",
    )
    document = spreadsheet(
        sheet("Rates", rates),
        sheet("Sales", sales),
    )

    result = compose({"Finance": document})

    formula = result.worksheets[1].tables[0].columns[0].excel_formula
    assert isinstance(formula, FormulaBinary)
    assert isinstance(formula.left, CellReference)
    assert formula.left.sheet_name == "Finance - Rates"
    assert result.worksheets[0].tables[0] is rates
    assert result.worksheets[1].tables[0].data.source is sales.data.source
    assert isinstance(amount.excel_formula, FormulaBinary)
    assert isinstance(amount.excel_formula.left, CellReference)
    assert amount.excel_formula.left.sheet_name == "Rates"


def test_compose_preserves_unqualified_references() -> None:
    formula = col("value") + table_ref("values").column("value").cell(0)
    rule = when(col("value") > 0, style=Style(fill="#D9EAF7"))
    values = table(
        source=[{"value": 1}],
        columns=(
            decimal(source="value"),
            decimal(id="total", formula=formula),
        ),
        name="values",
        rules=(rule,),
    )

    result = compose(
        {"Section": spreadsheet(sheet("Summary", values))},
    )

    assert result.worksheets[0].tables[0] is values
    assert values.columns[1].excel_formula is formula
    assert result.worksheets[0].tables[0].rules[0] is rule


def test_compose_rebases_qualified_range_references() -> None:
    formula = sheet_ref("Values").table("values").column("value")
    document = spreadsheet(
        sheet(
            "Values",
            table(
                source=[{"value": 1}],
                columns=(decimal(source="value"),),
                name="values",
            ),
        ),
        sheet(
            "Summary",
            table(
                source=[{}],
                columns=(decimal(id="total", formula=formula),),
                name="summary",
            ),
        ),
    )

    result = compose({"Section": document})

    composed_formula = result.worksheets[1].tables[0].columns[0].excel_formula
    assert isinstance(composed_formula, RangeReference)
    assert composed_formula.sheet_name == "Section - Values"
    assert formula.sheet_name == "Values"


def test_compose_preserves_unqualified_chart_source_identity() -> None:
    sales = table(
        source=[{"day": "Mon", "revenue": 10}],
        columns=(text(source="day"), decimal(source="revenue")),
        name="sales",
    )
    sales_chart = chart(table_ref("sales"), x="day", y="revenue")

    result = compose(
        {"Section": spreadsheet(sheet("Summary", sales, sales_chart))},
    )

    assert result.worksheets[0].blocks == (sales, sales_chart)
    assert result.worksheets[0].blocks[0] is sales
    assert result.worksheets[0].blocks[1] is sales_chart


def test_compose_preserves_inert_blocks_and_stack_by_identity() -> None:
    matrix_block = matrix(
        source=[],
        row="region",
        column="quarter",
        value=decimal(source="amount"),
    )
    title_block = title("Summary")
    spacer_block = spacer(rows=2)
    image_block = image(b"", name="logo")
    stack_block = stack(title("Nested"), spacer())
    blocks = (
        matrix_block,
        title_block,
        spacer_block,
        image_block,
        stack_block,
    )

    result = compose({"Section": spreadsheet(sheet("Summary", *blocks))})

    assert all(
        composed is original
        for composed, original in zip(
            result.worksheets[0].blocks,
            blocks,
            strict=True,
        )
    )


def test_compose_rejects_external_worksheet_reference_without_reading_rows() -> None:
    rows = _ObservableRows()
    document = spreadsheet(
        sheet(
            "Summary",
            table(
                source=rows,
                columns=(
                    decimal(
                        id="value",
                        formula=(
                            sheet_ref("External")
                            .table("values")
                            .column("value")
                            .cell(0)
                        ),
                    ),
                ),
                name="summary",
            ),
        ),
    )

    with pytest.raises(InvalidOperationError) as captured:
        compose({"Section": document})

    assert captured.value.context == {
        "reason": "cross_section_reference",
        "section": "Section",
        "worksheet": "Summary",
        "requested_worksheet": "External",
    }
    assert rows.iterations == 0


@pytest.mark.parametrize("backend", ["xlsxwriter", "openpyxl"])
def test_compose_and_validate_leave_one_shot_source_untouched(
    backend: str,
) -> None:
    iterations = 0

    def rows() -> Iterator[dict[str, int]]:
        nonlocal iterations
        iterations += 1
        yield {"value": 1}

    document = spreadsheet(
        sheet(
            "Summary",
            table(
                source=rows(),
                columns=(decimal(source="value"),),
                name="values",
            ),
        ),
    )

    composed = compose({"Section": document})
    validate(composed)

    assert iterations == 0
    render(composed, backend=backend)
    assert iterations == 1
    with pytest.raises(DataSourceConsumedError):
        render(composed, backend=backend)


def test_compose_rejects_an_unknown_formula_node() -> None:
    document = spreadsheet(
        sheet(
            "Summary",
            table(
                source=[{}],
                columns=(decimal(id="value", formula=_UnknownFormula()),),
            ),
        ),
    )

    with pytest.raises(UnsupportedFeatureError) as captured:
        compose({"Section": document})

    assert captured.value.path == "sections['Section'].worksheet['Summary']"
    assert captured.value.context == {
        "reason": "unsupported_formula_node",
        "section": "Section",
        "worksheet": "Summary",
        "formula_type": "_UnknownFormula",
    }


def test_compose_rejects_an_unknown_spreadsheet_block() -> None:
    document = spreadsheet(sheet("Summary"))
    object.__setattr__(  # noqa: PLC2801
        document.worksheets[0],
        "blocks",
        (object(),),
    )

    with pytest.raises(UnsupportedFeatureError) as captured:
        compose({"Section": document})

    assert captured.value.path == "sections['Section'].worksheet['Summary']"
    assert captured.value.context == {
        "reason": "unsupported_spreadsheet_block",
        "section": "Section",
        "worksheet": "Summary",
        "block_type": "object",
    }


def test_compose_inherits_named_styles_from_child_documents() -> None:
    heading = Style(fill="#D9EAF7")
    document = spreadsheet(
        sheet("Summary"),
        styles=StyleSheet({"heading": heading}),
    )

    result = compose({"Section": document})

    assert result.styles["heading"] is heading


def test_compose_coalesces_equal_styles_and_non_default_themes() -> None:
    first_heading = Style(fill="#D9EAF7")
    second_heading = Style(fill="#D9EAF7")
    first_theme = DocumentTheme(default=Style(fill="#FFFFFF"))
    second_theme = DocumentTheme(default=Style(fill="#FFFFFF"))
    first = spreadsheet(
        sheet("Summary"),
        styles={"heading": first_heading},
        theme=first_theme,
    )
    second = spreadsheet(
        sheet("Details"),
        styles={"heading": second_heading},
        theme=second_theme,
    )

    result = compose({"First": first, "Second": second})

    assert result.styles["heading"] is first_heading
    assert result.theme is first_theme


def test_compose_rejects_conflicting_child_styles() -> None:
    first = spreadsheet(
        sheet("Summary"),
        styles={"heading": Style(fill="#D9EAF7")},
    )
    second = spreadsheet(
        sheet("Summary"),
        styles={"heading": Style(fill="#FCE8E6")},
    )

    with pytest.raises(InvalidOperationError) as captured:
        compose({"First": first, "Second": second})

    assert captured.value.context == {
        "reason": "style_conflict",
        "style": "heading",
        "section": "Second",
        "conflicts_with": "First",
    }


def test_compose_inherits_one_non_default_child_theme() -> None:
    report_theme = DocumentTheme(default=Style(fill="#FFFFFF"))
    default_document = spreadsheet(sheet("Summary"))
    themed_document = spreadsheet(sheet("Details"), theme=report_theme)

    result = compose(
        {
            "Default": default_document,
            "Themed": themed_document,
        },
    )

    assert result.theme is report_theme


def test_compose_merges_non_conflicting_child_metadata() -> None:
    params = spreadsheet(
        sheet("Summary"),
        metadata={"owner": "security"},
    )
    vulnerabilities = spreadsheet(
        sheet("Details"),
        metadata={"period": "2026-Q3"},
    )

    result = compose(
        {
            "Params": params,
            "Vulnerabilities": vulnerabilities,
        },
    )

    assert result.metadata == {
        "owner": "security",
        "period": "2026-Q3",
    }


def test_compose_treats_nested_metadata_as_opaque_values() -> None:
    first = spreadsheet(
        sheet("Summary"),
        metadata={"options": {"locale": "en"}},
    )
    equal = spreadsheet(
        sheet("Details"),
        metadata={"options": {"locale": "en"}},
    )

    result = compose({"First": first, "Equal": equal})

    assert result.metadata["options"] == {"locale": "en"}

    conflicting = spreadsheet(
        sheet("Other"),
        metadata={"options": {"timezone": "UTC"}},
    )
    with pytest.raises(InvalidOperationError) as captured:
        compose({"First": first, "Conflicting": conflicting})

    assert captured.value.context == {
        "reason": "metadata_conflict",
        "key": "options",
        "section": "Conflicting",
        "conflicts_with": "First",
    }


def test_compose_rejects_template_bearing_child_document() -> None:
    document = spreadsheet(
        sheet("Summary"),
        template=TemplateSpecification(
            source="template.xlsx",
            format="xlsx",
        ),
    )

    with pytest.raises(InvalidOperationError) as captured:
        compose({"Template": document})

    assert captured.value.context == {
        "reason": "included_template",
        "section": "Template",
    }


def test_compose_rebases_nested_chart_and_conditional_rule_references() -> None:
    rate = sheet_ref("Rates").table("rates").column("rate").cell(0)
    rates = table(
        source=[{"rate": 2}],
        columns=(decimal(source="rate"),),
        name="rates",
    )
    sales = table(
        source=[{"day": "Mon", "revenue": 10}],
        columns=(
            text(source="day"),
            decimal(source="revenue"),
        ),
        name="sales",
        rules=(when(rate > 0, style=Style(fill="#D9EAF7")),),
    )
    dashboard = stack(
        chart(
            sheet_ref("Sales").table("sales"),
            x="day",
            y="revenue",
        ),
    )
    document = spreadsheet(
        sheet("Rates", rates),
        sheet("Sales", sales),
        sheet("Dashboard", dashboard),
    )

    result = compose({"Finance": document})

    condition = result.worksheets[1].tables[0].rules[0].condition
    assert isinstance(condition, FormulaBinary)
    assert isinstance(condition.left, CellReference)
    assert condition.left.sheet_name == "Finance - Rates"
    composed_stack = result.worksheets[2].blocks[0]
    assert isinstance(composed_stack, Stack)
    composed_chart = composed_stack.items[0]
    assert isinstance(composed_chart, Chart)
    assert composed_chart.source.sheet_name == "Finance - Sales"


def test_rendered_composed_chart_uses_rebased_source_ranges() -> None:
    document = spreadsheet(
        sheet(
            "Sales",
            table(
                source=[
                    {"day": "Mon", "revenue": 10},
                    {"day": "Tue", "revenue": 20},
                ],
                columns=(text(source="day"), decimal(source="revenue")),
                name="sales",
            ),
        ),
        sheet(
            "Dashboard",
            chart(
                sheet_ref("Sales").table("sales"),
                x="day",
                y="revenue",
            ),
        ),
    )

    result = render(compose({"Finance": document}), backend="xlsxwriter")
    workbook = load_workbook(BytesIO(result.data or b""))
    rendered_charts = getattr(  # noqa: B009
        workbook["Finance - Dashboard"],
        "_charts",
    )
    series = rendered_charts[0].series[0]

    assert series.cat is not None
    assert series.cat.strRef is not None
    assert series.cat.strRef.f == "'Finance - Sales'!$A$2:$A$3"
    assert series.val is not None
    assert series.val.numRef is not None
    assert series.val.numRef.f == "'Finance - Sales'!$B$2:$B$3"


def test_compose_uses_outer_style_to_resolve_child_conflict() -> None:
    body = Style(fill="#EEEEEE")
    first = spreadsheet(
        sheet("Summary"),
        styles={"body": body, "heading": Style(fill="#D9EAF7")},
    )
    second = spreadsheet(
        sheet("Summary"),
        styles={"heading": Style(fill="#FCE8E6")},
    )
    heading = Style(fill="#FFFFFF")

    result = compose(
        {"First": first, "Second": second},
        styles=StyleSheet({"heading": heading}),
    )

    assert result.styles["heading"] is heading
    assert result.styles["body"] is body


def test_compose_requires_outer_theme_for_conflicting_children() -> None:
    first = spreadsheet(
        sheet("Summary"),
        theme=DocumentTheme(default=Style(fill="#D9EAF7")),
    )
    second = spreadsheet(
        sheet("Summary"),
        theme=DocumentTheme(default=Style(fill="#FCE8E6")),
    )

    with pytest.raises(InvalidOperationError) as captured:
        compose({"First": first, "Second": second})

    assert captured.value.context == {
        "reason": "theme_conflict",
        "section": "Second",
        "conflicts_with": "First",
    }

    selected = DocumentTheme(default=Style(fill="#FFFFFF"))
    result = compose(
        {"First": first, "Second": second},
        theme=selected,
    )
    assert result.theme is selected


def test_compose_requires_outer_metadata_for_conflicting_children() -> None:
    first = spreadsheet(
        sheet("Summary"),
        metadata={"owner": "first", "period": "2026-Q3"},
    )
    second = spreadsheet(sheet("Summary"), metadata={"owner": "second"})

    with pytest.raises(InvalidOperationError) as captured:
        compose({"First": first, "Second": second})

    assert captured.value.context == {
        "reason": "metadata_conflict",
        "key": "owner",
        "section": "Second",
        "conflicts_with": "First",
    }

    result = compose(
        {"First": first, "Second": second},
        metadata={"owner": "combined"},
    )
    assert result.metadata == {"owner": "combined"}


def test_compose_preserves_outer_template() -> None:
    specification = TemplateSpecification(
        source="template.xlsx",
        format="xlsx",
    )

    result = compose(
        {"Section": spreadsheet(sheet("Summary"))},
        template=specification,
    )

    assert result.template is specification


def test_compose_leaves_global_table_conflicts_to_validation() -> None:
    first = spreadsheet(
        sheet(
            "Summary",
            table(
                source=[],
                columns=(text(source="value"),),
                name="issues",
            ),
        ),
    )
    second = spreadsheet(
        sheet(
            "Summary",
            table(
                source=[],
                columns=(text(source="value"),),
                name="issues",
            ),
        ),
    )

    with pytest.raises(ValidationError, match="Duplicate table name") as captured:
        validate(compose({"First": first, "Second": second}))

    duplicate = next(
        issue for issue in captured.value.issues if issue.code == "duplicate_table"
    )
    assert duplicate.context == {
        "table": "issues",
        "worksheet": "Second - Summary",
        "conflicts_with_table": "issues",
        "conflicts_with_worksheet": "First - Summary",
    }


@pytest.mark.parametrize("backend", ["xlsxwriter", "openpyxl"])
def test_compose_renders_each_source_once(backend: str) -> None:
    value_rows = _ObservableRows()
    summary_rows = _ObservableRows()
    document = spreadsheet(
        sheet(
            "Values",
            table(
                source=value_rows,
                columns=(decimal(source="value"),),
                name="values",
            ),
        ),
        sheet(
            "Summary",
            table(
                source=summary_rows,
                columns=(
                    decimal(
                        id="first_value",
                        formula=(
                            sheet_ref("Values")
                            .table("values")
                            .column("value")
                            .cell(0)
                            .absolute()
                        ),
                    ),
                ),
                name="summary",
            ),
        ),
    )

    composed = compose({"Section": document})
    assert value_rows.iterations == 0
    assert summary_rows.iterations == 0

    artifact = inspect_artifact(render(composed, backend=backend))

    assert artifact.worksheet("Section - Values").cell("A2").value == 1
    assert artifact.worksheet("Section - Summary").cell("A2").formula == (
        "='Section - Values'!$A$2"
    )
    assert value_rows.iterations == 1
    assert summary_rows.iterations == 1
